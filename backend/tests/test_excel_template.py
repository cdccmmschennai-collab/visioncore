"""Tests for Template workbook display-only formatting:
- Additional Information: each "LABEL: value" pair on its own line within
  the same cell.
- Size / Dimension: the inch mark (") spelled out as "inch".
Both are Template workbook only — the AI Output workbook, the database and
the UI all keep the original raw string untouched.
"""
from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_template import (
    _format_additional_information,
    _format_size_dimension,
    build_template_workbook,
)
from app.services.fields import empty_payload, value_of


def test_format_additional_information_breaks_each_pair_onto_its_own_line():
    raw = (
        "TYPE/RATING: 1500# RTJ X 6000# NPT F, BODY MOC: ASTM A350 GR LF2 CL.1, "
        "TRIM MOC: ASTM A182 GR F316/F316L"
    )
    assert _format_additional_information(raw) == (
        "TYPE/RATING: 1500# RTJ X 6000# NPT F\n"
        "BODY MOC: ASTM A350 GR LF2 CL.1\n"
        "TRIM MOC: ASTM A182 GR F316/F316L"
    )


def test_format_additional_information_leaves_a_single_pair_unchanged():
    assert _format_additional_information("TYPE/RATING: 1500# RTJ") == "TYPE/RATING: 1500# RTJ"


def test_format_additional_information_ignores_a_comma_that_is_not_a_pair_boundary():
    """A comma inside a value's own text (not immediately followed by a new
    "LABEL:") must not be split — only real pair boundaries are."""
    raw = "STANDARD: ASME B16.5, API 600 combined rating"
    assert _format_additional_information(raw) == raw


def test_build_template_workbook_writes_line_breaks_and_keeps_wrap_text():
    payload = empty_payload("12-TAG-0001", "BALL VALVE")
    payload["fields"]["additional_information"] = {
        "value": (
            "TYPE/RATING: 1500# RTJ X 6000# NPT F, BODY MOC: ASTM A350 GR LF2 CL.1, "
            "TRIM MOC: ASTM A182 GR F316/F316L"
        ),
        "quality": "Confirmed",
    }

    content = build_template_workbook([{"payload": payload, "ai_payload": {}}])
    wb = load_workbook(BytesIO(content))
    ws = wb.active

    header_row = [c.value for c in ws[1]]
    col = header_row.index("ADDITIONAL INFORMATION") + 1
    cell = ws.cell(row=2, column=col)

    assert cell.value == (
        "TYPE/RATING: 1500# RTJ X 6000# NPT F\n"
        "BODY MOC: ASTM A350 GR LF2 CL.1\n"
        "TRIM MOC: ASTM A182 GR F316/F316L"
    )
    assert cell.alignment.wrap_text is True


def test_format_size_dimension_spells_out_the_inch_mark():
    assert _format_size_dimension('6"') == "6 inch"
    assert _format_size_dimension('12"') == "12 inch"
    assert _format_size_dimension('24" x 6"') == "24 inch x 6 inch"


def test_format_size_dimension_leaves_a_value_without_the_mark_unchanged():
    assert _format_size_dimension("3/4 150RF") == "3/4 150RF"


def test_build_template_workbook_spells_out_inch_mark_in_size_column_only():
    payload = empty_payload("12-TAG-0001", "BALL VALVE")
    payload["fields"]["size_dimension"] = {"value": '24" x 6"', "quality": "Confirmed"}
    ai_payload = {"fields": {"size_dimension": {"value": '24" x 6"', "quality": "Confirmed"}}}

    content = build_template_workbook([{"payload": payload, "ai_payload": ai_payload}])
    wb = load_workbook(BytesIO(content))
    ws = wb.active

    header_row = [c.value for c in ws[1]]
    col = header_row.index("SIZE/DIMENSION") + 1
    assert ws.cell(row=2, column=col).value == "24 inch x 6 inch"

    # The underlying payload the database/UI/AI Output workbook see is untouched.
    assert value_of(payload, "size_dimension") == '24" x 6"'


def _asset_tag_row(input_filename, ai_payload):
    payload = empty_payload("22-4203-BV-0105", "BALL VALVE")
    content = build_template_workbook([{
        "payload": payload, "ai_payload": ai_payload, "input_filename": input_filename,
    }])
    ws = load_workbook(BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    tag_col = headers.index("TAG NUMBER") + 1
    assert headers[tag_col] == "ASSET TAG NUMBER"
    return ws.cell(row=2, column=tag_col), ws.cell(row=2, column=tag_col + 1)


def test_tag_number_comes_from_filename_and_asset_tag_from_photo():
    tag, asset = _asset_tag_row(
        "22-4203-BV-0105-BALL VALVE.jpg", {"photo_tag_number": "22-4203-BV-0105"}
    )
    assert tag.value == "22-4203-BV-0105"
    assert asset.value == "22-4203-BV-0105"
    assert asset.fill.fgColor.rgb != "FFFFA500"


def test_tag_only_filename_parses():
    tag, _ = _asset_tag_row("22-4202-TW-0021.jpg", {"photo_tag_number": ""})
    assert tag.value == "22-4202-TW-0021"


def test_tag_number_drops_comma_description_from_filename():
    tag, _ = _asset_tag_row("76-FT-301-TRANSMITTER,FLOW.jpg", {"photo_tag_number": "76-FT-301"})
    assert tag.value == "76-FT-301"


def test_ai_workbook_tag_number_drops_comma_description_from_filename():
    from app.services.excel_ai import build_ai_workbook

    payload = empty_payload("76-FT-301-A", "TRANSMITTER,FLOW")
    content = build_ai_workbook(
        payload, "76-FT-301-A", "TRANSMITTER,FLOW", ["76-FT-301-TRANSMITTER,FLOW.jpg"]
    )
    ws = load_workbook(BytesIO(content)).active
    assert ws["B5"].value == "76-FT-301"


def test_asset_tag_mismatch_is_orange():
    _, asset = _asset_tag_row(
        "22-4203-BV-0105-BALL VALVE.jpg", {"photo_tag_number": "22-4203-BV-0106"}
    )
    assert asset.value == "22-4203-BV-0106"
    assert asset.fill.fgColor.rgb == "FFFFA500"


def test_asset_tag_empty_when_photo_had_none_and_no_filename_fallback():
    _, asset = _asset_tag_row(
        "22-4203-BV-0105-BALL VALVE.jpg", {"photo_tag_number": ""}
    )
    assert asset.value is None
    assert asset.fill.fgColor.rgb != "FFFFA500"


def test_autofilter_spans_every_column():
    content = build_template_workbook([{"payload": empty_payload("T-1", "X"), "ai_payload": {}}])
    ws = load_workbook(BytesIO(content)).active
    assert ws.auto_filter.ref == "A1:R2"
    assert ws.max_column == 18


def test_template_omits_description_column():
    content = build_template_workbook([{"payload": empty_payload("T-1", "BALL VALVE"), "ai_payload": {}}])
    ws = load_workbook(BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    assert "EQUIPMENT DESCRIPTION" not in headers
    assert headers[1:4] == ["TAG NUMBER", "ASSET TAG NUMBER", "SIZE/DIMENSION"]
    assert "BALL VALVE" not in [c.value for c in ws[2]]


def test_tag_number_never_amber_in_template():
    payload = empty_payload("22-4203-BV-0106", "BALL VALVE")
    payload["fields"]["tag_number"]["quality"] = "Verify"
    content = build_template_workbook([{
        "payload": payload, "ai_payload": {},
        "input_filename": "22-4203-BV-0105-BALL VALVE.jpg",
    }])
    ws = load_workbook(BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    cell = ws.cell(row=2, column=headers.index("TAG NUMBER") + 1)
    assert cell.value == "22-4203-BV-0105"
    assert cell.fill.fgColor.rgb != "FFFFC000"


def test_ai_workbook_tag_number_from_filename_and_confirmed():
    from app.services.excel_ai import build_ai_workbook

    payload = empty_payload("22-4203-BV-0106", "BALL VALVE")
    payload["fields"]["tag_number"]["quality"] = "Verify"
    content = build_ai_workbook(
        payload, "22-4203-BV-0106", "BALL VALVE", ["22-4203-BV-0105-BALL VALVE.jpg"]
    )
    ws = load_workbook(BytesIO(content)).active
    assert ws["A5"].value == "Tag Number"
    assert ws["B5"].value == "22-4203-BV-0105"
    assert ws["C5"].value == "Confirmed"


def test_additional_information_paragraph_column():
    payload = empty_payload("12-TAG-0001", "BALL VALVE")
    raw = "TYPE/RATING: 1500# RTJ, BODY MOC: ASTM A350 GR LF2 CL.1, TRIM MOC: F316"
    payload["fields"]["additional_information"] = {"value": raw, "quality": "Verify"}
    content = build_template_workbook([{"payload": payload, "ai_payload": {}}])
    ws = load_workbook(BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    col = headers.index("ADDITIONAL INFORMATION") + 1
    assert headers[col] == "ADDITIONAL INFORMATION-PARAGRAPH"
    assert headers[col + 1] == "REMARKS"
    listed, para = ws.cell(row=2, column=col), ws.cell(row=2, column=col + 1)
    assert "\n" in listed.value
    assert para.value == raw
    assert para.fill.fgColor.rgb == listed.fill.fgColor.rgb
    assert para.alignment.wrap_text is True


def test_additional_information_paragraph_empty_when_not_present():
    content = build_template_workbook([{"payload": empty_payload("T-1", "X"), "ai_payload": {}}])
    ws = load_workbook(BytesIO(content)).active
    headers = [c.value for c in ws[1]]
    col = headers.index("ADDITIONAL INFORMATION-PARAGRAPH") + 1
    assert ws.cell(row=2, column=col).value is None


def test_ai_workbook_already_extracted_banner_only_when_requested():
    from app.services.excel_ai import ALREADY_EXTRACTED_NOTE, build_ai_workbook

    payload = empty_payload("22-4203-BV-0105", "BALL VALVE")
    plain = load_workbook(BytesIO(build_ai_workbook(payload, "22-4203-BV-0105", "BALL VALVE", []))).active
    assert plain["A3"].value is None

    marked = load_workbook(BytesIO(build_ai_workbook(
        payload, "22-4203-BV-0105", "BALL VALVE", [], status_note=ALREADY_EXTRACTED_NOTE,
    ))).active
    assert marked["A3"].value.startswith("ALREADY EXTRACTED")
    # Table layout below is unchanged.
    assert marked["A4"].value == "Field"
    assert marked["A5"].value == "Tag Number"


def test_duplicate_highlights_only_its_tag_number_cell():
    records = [
        {"payload": empty_payload("T-1", "X"), "ai_payload": {}, "duplicate": True},
        {"payload": empty_payload("T-2", "X"), "ai_payload": {}},
    ]
    ws = load_workbook(BytesIO(build_template_workbook(records))).active
    headers = [c.value for c in ws[1]]
    tag_col = headers.index("TAG NUMBER") + 1
    assert ws.cell(2, tag_col).fill.fgColor.rgb == "FFC6EFCE"
    others = [ws.cell(2, c).fill.fgColor.rgb for c in range(1, len(headers) + 1) if c != tag_col]
    assert "FFC6EFCE" not in others
    assert ws.cell(3, tag_col).fill.fgColor.rgb != "FFC6EFCE"
